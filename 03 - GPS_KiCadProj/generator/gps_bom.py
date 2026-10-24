# ACwing GPS rev A: assembly BoM (+ Costing sheet) and JLCPCB BoM.
#   python3 gps_bom.py <out dir>
import sys, os, re, json
import pcbnew
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

OUT = sys.argv[1]
PROJECT = 'ACwing_GPS_A'
parts = json.load(open(os.path.join(OUT, PROJECT + '.kicad_sch.parts.json')))
b = pcbnew.LoadBoard(os.path.join(OUT, PROJECT + '.kicad_pcb'))
side = {f.GetReference(): ('Bottom' if f.IsFlipped() else 'Top') for f in b.GetFootprints()}
key = lambda r: (re.sub(r'\d', '', r), int(re.sub(r'\D', '', r) or 0))
Q, E = 'LCSC quote', 'Estimate'
# (value, footprint) -> MPN, manufacturer, LCSC/JLC #, package, description, note, JLC footprint label, unit price
#                       10 boards, unit price 100 boards (USD), price source
TAB = {
 ('SAM-M10Q-00B', 'ublox_SAM-M8Q'): ('SAM-M10Q-00B', 'u-blox', 'Consigned (not at LCSC)', 'LGA-20 15.5 x 15.5 x 6.3 mm',
     'GNSS antenna module, u-blox M10, GPS/Galileo/BeiDou/GLONASS, UART + I2C', 'Buy from DigiKey/Mouser and send to JLC as consigned parts, or solder by hand. MSL 4: bake before reflow if the bag was opened',
     'LGA-20 15.5x15.5', 22.0, 19.0, E),
 ('IIS2MDC', 'LGA-12_2x2mm_P0.5mm'): ('IIS2MDCTR', 'STMicroelectronics', 'C2655002', 'LGA-12 (2x2)', '3-axis magnetometer, I2C 0x1E', 'Axes aligned with the FC; same part as the FC compass', 'LGA-12 (2x2)', 3.69, 2.98, Q),
 ('Qwiic IN', 'JST_SH_SM04B-SRSS-TB_1x04-1MP_P1.00mm_Horizontal'): ('SM04B-SRSS-TB(LF)(SN)', 'JST', 'C160404', 'JST SH 4P side entry, h 2.9 mm',
     'Connector, JST SH, 4 positions, 1.0 mm pitch, side entry', 'Qwiic in: top side, right edge at v = +3.6 (aligned with FC J3), mouth +X; cable from FC J3', 'SMD,P=1mm,Side entry (JST SH 4P)', 0.25, 0.18, E),
 ('Qwiic OUT', 'JST_SH_SM04B-SRSS-TB_1x04-1MP_P1.00mm_Horizontal'): ('SM04B-SRSS-TB(LF)(SN)', 'JST', 'C160404', 'JST SH 4P side entry, h 2.9 mm',
     'Connector, JST SH, 4 positions, 1.0 mm pitch, side entry', 'Qwiic out: top side, mouth at the rear edge; optional pass-through', 'SMD,P=1mm,Side entry (JST SH 4P)', 0.25, 0.18, E),
 ('10uF', 'C_0603_1608Metric'): ('CL10A106KP8NNNC', 'Samsung Electro-Mechanics', 'C19702', '0603 (1608 metric), h 0.9 mm', 'Capacitor, MLCC, 10 uF, 10 V, X5R', '', '0603', 0.0287, 0.0287, Q),
 ('100nF', 'C_0402_1005Metric'): ('CL05B104KO5NNNC', 'Samsung Electro-Mechanics', 'C1525', '0402 (1005 metric)', 'Capacitor, MLCC, 100 nF, 16 V, X7R', '', '0402', 0.0049, 0.0049, Q),
 ('220nF', 'C_0402_1005Metric'): ('CL05B224KO5NNNC', 'Samsung Electro-Mechanics', 'C16772', '0402 (1005 metric)', 'Capacitor, MLCC, 220 nF, 16 V, X7R', 'IIS2MDC C1 (ESR < 200 mOhm)', '0402', 0.0054, 0.0054, Q),
 ('11mF', 'C_1210_3225Metric'): ('CPH3225A', 'Seiko Instruments', 'Consigned (not at LCSC)', '3225 (1210), h 0.9 mm',
     'Electric double-layer capacitor, 11 mF, 3.3 V, reflow capable', 'GNSS backup supply (hot start). Check the land pattern against the Seiko drawing; + toward D1', '1210', 1.30, 1.00, E),
 ('RB521S30', 'D_SOD-523'): ('RB521S30', 'Changjiang (JCET)', 'C8523', 'SOD-523', 'Diode, Schottky, 30 V, 200 mA', 'Backup supply blocking diode', 'SOD-523', 0.012, 0.008, E),
 ('2N7002', 'SOT-23'): ('2N7002', 'Changjiang (JCET)', 'C8545', 'SOT-23', 'MOSFET, N-channel, 60 V, 115 mA', 'PPS LED driver, gate on TIMEPULSE only', 'SOT-23', 0.012, 0.008, E),
 ('green', 'LED_0402_1005Metric'): ('XL-1005UGC', 'XINGLIGHT', 'C965793', '0402 (1005 metric)', 'LED, green, 0402', 'PPS indicator, top side', '0402', 0.012, 0.008, E),
 ('1k', 'R_0402_1005Metric'): ('0402WGF1001TCE', 'UNI-ROYAL', 'C11702', '0402 (1005 metric)', 'Resistor, 1 kOhm, 1 %, 1/16 W', 'Supercap charge limit', '0402', 0.001, 0.001, E),
 ('220R', 'R_0402_1005Metric'): ('0402WGF2200TCE', 'UNI-ROYAL', 'C25091', '0402 (1005 metric)', 'Resistor, 220 Ohm, 1 %, 1/16 W', 'LED current ~1.4 mA', '0402', 0.001, 0.001, E),
}
NOT_ASSEMBLED = [
 ('H', 'M2', 'N/A (PCB feature)', 'N/A', 'NPTH 2.2 mm', 'Mounting hole, M2, 20 x 20 pattern', 'Keep-out r 2.6 mm on both sides for the nylon standoff / nut'),
 ('TP', 'Test pad', 'N/A (PCB feature)', 'N/A', 'SMD pad 1.0 mm', 'Test pads SAFEBOOT_N, RESET_N, GND, GNSS TXD, GNSS RXD', 'Firmware recovery: short TP1 to TP3 at power-up; u-center over a 3.3 V USB-UART on TP4/TP5'),
]
groups, skipped, dnp = {}, [], []
for ref, v in parts.items():
    fpn = v['fp'].split(':')[1]
    if ref.startswith(('TP', 'H')):
        skipped.append(ref); continue
    if v.get('dnp'):
        dnp.append(ref); continue
    k = (v['value'], fpn)
    if k not in TAB:
        raise SystemExit('no MPN for %s %s' % (ref, k))
    groups.setdefault((k, side[ref]), []).append(ref)

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
def style_header(ws, r):
    for c in ws[r]:
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
wb = Workbook(); ws = wb.active; ws.title = 'BoM'
hdr = ['Item', 'Quantity', 'Part Reference', 'Value', 'MPN', 'Manufacturer', 'LCSC / JLCPCB #', 'Package', 'Description', 'Side', 'Notes']
ws.append(['ACwing GPS rev A - assembly BoM (1 board)'])
ws.append(['Source: ACwing_GPS_A.kicad_sch / ACwing_GPS_A.kicad_pcb. LCSC part numbers checked on 2026-10-17. U1 and C5 are not stocked by LCSC: consigned parts or hand soldering.'])
ws.append([]); ws.append(hdr); H = 4
rows = sorted(groups.items(), key=lambda kv: key(sorted(kv[1], key=key)[0]))
for i, ((k, sd), refs) in enumerate(rows, 1):
    t = TAB[k]; refs = sorted(refs, key=key)
    ws.append([i, len(refs), ', '.join(refs), k[0], t[0], t[1], t[2], t[3], t[4], sd, t[5]])
last = ws.max_row
ws.append([])
ws.append(['Total', '=SUM(B%d:B%d)' % (H + 1, last), 'Total placements (assembled parts)', '-', '=COUNTA(E%d:E%d)' % (H + 1, last),
           'line items', '-', '-', '-', '-', 'Not-assembled items listed below'])
tot = ws.max_row
ws.append([]); ws.append(['Not assembled (PCB features / DNP)']); na0 = ws.max_row
ws.append(hdr); NA_H = ws.max_row
n = 0
for pref, val, mpn, man, pkg, desc, note in NOT_ASSEMBLED:
    refs = sorted([r for r in skipped if re.match(pref + r'\d', r)], key=key)
    n += 1
    ws.append(['NA%d' % n, len(refs), ', '.join(refs), val, mpn, man, 'N/A', pkg, desc, ', '.join(sorted({side[r] for r in refs})), note])
for r in sorted(dnp, key=key):
    n += 1
    ws.append(['NA%d' % n, 1, r, parts[r]['value'], '0402WGF4701TCE', 'UNI-ROYAL', 'C25900 (DNP)', '0402 (1005 metric)',
               'Resistor, 4.7 kOhm, 1 %, 1/16 W', side[r], 'Do not populate: the I2C pull-ups are on the FC (R15/R16). Fit only for a host without pull-ups'])
end = ws.max_row
for row in ws.iter_rows(min_row=1, max_row=end):
    for c in row:
        c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
ws['A1'].font = Font(name=FONT, size=13, bold=True)
ws['A2'].font = Font(name=FONT, size=9, italic=True, color='595959')
style_header(ws, H); style_header(ws, NA_H)
ws.cell(na0, 1).font = Font(name=FONT, size=11, bold=True)
for r in list(range(H + 1, last + 1)) + list(range(NA_H + 1, end + 1)):
    for c in ws[r]:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for c in ws[tot]:
    c.font = Font(name=FONT, size=10, bold=True)
for col, w in zip('ABCDEFGHIJK', (6, 9, 22, 14, 24, 22, 20, 24, 40, 9, 60)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'

# ---------------------------------------------------------------- Costing
cs = wb.create_sheet('Costing')
cs.append(['ACwing GPS rev A - component costing (assembled parts, USD)'])
cs.append(['Blue = input prices; black = formulas. "Estimate" rows are typical distributor prices, not quotes. PCB, assembly fees and the cable are not included.'])
cs.append([])
ch = ['References', 'Qty / board', 'MPN', 'Source', 'Unit price @10 boards', 'Unit price @100 boards', 'Cost / board @10', 'Cost / board @100', 'Price basis']
cs.append(ch); CH = 4
for (k, sd), refs in rows:
    t = TAB[k]; r = cs.max_row + 1
    cs.append([', '.join(sorted(refs, key=key)), len(refs), t[0], t[2], t[7], t[8], '=B%d*E%d' % (r, r), '=B%d*F%d' % (r, r), t[9]])
cl = cs.max_row
cs.append([])
cs.append(['Total per board', '=SUM(B%d:B%d)' % (CH + 1, cl), '', '', '', '', '=SUM(G%d:G%d)' % (CH + 1, cl), '=SUM(H%d:H%d)' % (CH + 1, cl), ''])
ct = cs.max_row
u1 = next(r for r in range(CH + 1, cl + 1) if cs.cell(r, 1).value == 'U1')
cs.append(['Share of U1 (SAM-M10Q)', '', '', '', '', '', '=G%d/G%d' % (u1, ct), '=H%d/H%d' % (u1, ct), 'The GNSS module is most of the cost'])
for row in cs.iter_rows(min_row=1, max_row=cs.max_row):
    for c in row:
        c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
cs['A1'].font = Font(name=FONT, size=13, bold=True)
cs['A2'].font = Font(name=FONT, size=9, italic=True, color='595959')
style_header(cs, CH)
for r in range(CH + 1, cl + 1):
    for col in (5, 6):
        cs.cell(r, col).font = Font(name=FONT, size=10, color='0000FF'); cs.cell(r, col).number_format = '$#,##0.000'
    for col in (7, 8):
        cs.cell(r, col).number_format = '$#,##0.000'
    for c in cs[r]:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for col in (7, 8):
    cs.cell(ct, col).number_format = '$#,##0.00'; cs.cell(ct + 1, col).number_format = '0.0%'
for c in cs[ct]:
    c.font = Font(name=FONT, size=10, bold=True)
for col, w in zip('ABCDEFGHI', (22, 10, 24, 24, 13, 13, 14, 14, 16)):
    cs.column_dimensions[col].width = w
wb.save(os.path.join(OUT, PROJECT + '_BoM.xlsx'))

# ---------------------------------------------------------------- JLCPCB BoM
jw = Workbook(); js = jw.active; js.title = 'Sheet1'
js.append(['Comment', 'Designator', 'Footprint', 'JLCPCB Part #'])
jg = {}
for (k, sd), refs in rows:
    t = TAB[k]
    g = jg.setdefault(t[0], [t, []]); g[1] += refs
for mpn, (t, refs) in sorted(jg.items(), key=lambda kv: key(sorted(kv[1][1], key=key)[0])):
    js.append([mpn, ','.join(sorted(refs, key=key)), t[6], t[2] if t[2].startswith('C') and t[2][1:].isdigit() else ''])
for row in js.iter_rows():
    for c in row:
        c.font = Font(name=FONT, size=10, bold=(c.row == 1))
for col, w in zip('ABCD', (26, 30, 34, 16)):
    js.column_dimensions[col].width = w
jw.save(os.path.join(OUT, PROJECT + '_BoM_JLCPCB.xlsx'))
print('BoM lines', len(rows), 'placements', sum(len(v) for v in groups.values()), '| JLC lines', len(jg), '| not assembled', n)
